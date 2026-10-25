#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Universal file reader for WLstudio read-binary plugin."""

import sys, os, json, zipfile, re
from pathlib import Path

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

try:
    import PyPDF2
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

try:
    from striprtf.striprtf import rtf_to_text
    HAS_RTF = True
except ImportError:
    HAS_RTF = False

try:
    import ebooklib
    from ebooklib import epub
    from bs4 import BeautifulSoup
    HAS_EPUB = True
except ImportError:
    HAS_EPUB = False

try:
    import olefile
    HAS_OLE = True
except ImportError:
    HAS_OLE = False

try:
    import xlrd
    HAS_XLRD = True
except ImportError:
    HAS_XLRD = False


def detect_format(file_path):
    ext = Path(file_path).suffix.lower()
    if ext in ['.docx','.xlsx','.pptx','.odt','.ods','.odp','.wps','.et','.dps']:
        try:
            with zipfile.ZipFile(file_path, 'r') as z:
                names = z.namelist()
                if 'word/document.xml' in names: return 'docx'
                elif 'xl/workbook.xml' in names: return 'xlsx'
                elif 'ppt/presentation.xml' in names: return 'pptx'
                elif 'content.xml' in names:
                    if ext == '.ods': return 'ods'
                    elif ext == '.odp': return 'odp'
                    else: return 'odt'
        except:
            pass
    if ext in ['.doc','.xls','.ppt','.wps','.et','.dps']:
        try:
            if HAS_OLE and olefile.isOleFile(file_path):
                return 'ole2'
        except:
            pass
    ext_map = {
        '.docx':'docx','.xlsx':'xlsx','.pptx':'pptx',
        '.wps':'wps','.et':'et','.dps':'dps',
        '.pdf':'pdf','.rtf':'rtf','.epub':'epub',
        '.odt':'odt','.ods':'ods','.odp':'odp',
        '.doc':'doc','.xls':'xls','.ppt':'ppt',
        '.zip':'zip','.rar':'archive','.7z':'archive',
        '.tar':'archive','.gz':'archive','.bz2':'archive',
    }
    return ext_map.get(ext, 'unknown')


def read_docx(file_path):
    if not HAS_DOCX: return None, 'python-docx not installed'
    try:
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        tables_text = []
        for table in doc.tables:
            rows = [' | '.join([cell.text for cell in row.cells]) for row in table.rows]
            if rows: tables_text.append('\n'.join(rows))
        result = '\n\n'.join(paragraphs)
        if tables_text:
            result += '\n\n[表格内容]\n' + '\n\n[表格]\n'.join(tables_text)
        return result, None
    except Exception as e:
        return None, f'docx error: {e}'


def read_xlsx(file_path):
    if not HAS_OPENPYXL: return None, 'openpyxl not installed'
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        parts = []
        for name in wb.sheetnames:
            sheet = wb[name]
            rows = []
            for row in sheet.iter_rows():
                vals = [str(c.value) if c.value is not None else '' for c in row]
                if any(v.strip() for v in vals): rows.append(' | '.join(vals))
            parts.append(f'=== 工作表: {name} ({len(rows)} 行) ===\n' + '\n'.join(rows))
        return '\n\n'.join(parts), None
    except Exception as e:
        return None, f'xlsx error: {e}'


def read_pptx(file_path):
    if not HAS_PPTX: return None, 'python-pptx not installed'
    try:
        prs = Presentation(file_path)
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [s.text for s in slide.shapes if hasattr(s,'text') and s.text.strip()]
            if texts: parts.append(f'--- 幻灯片 {i} ---\n' + '\n'.join(texts))
        return '\n\n'.join(parts) if parts else '[未提取到文本]', None
    except Exception as e:
        return None, f'pptx error: {e}'


def read_pdf(file_path):
    if not HAS_PDF: return None, 'PyPDF2 not installed'
    try:
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            parts = []
            for i, page in enumerate(reader.pages, 1):
                text = page.extract_text()
                if text and text.strip(): parts.append(f'--- 第 {i} 页 ---\n{text}')
            meta = []
            md = reader.metadata
            if md:
                if md.title: meta.append(f'标题: {md.title}')
                if md.author: meta.append(f'作者: {md.author}')
                if md.subject: meta.append(f'主题: {md.subject}')
            result = ''
            if meta: result = '[PDF 元信息]\n' + '\n'.join(meta) + '\n\n'
            result += '\n\n'.join(parts) if parts else '[此 PDF 未提取到文本]'
            return result, None
    except Exception as e:
        return None, f'pdf error: {e}'


def read_rtf(file_path):
    if not HAS_RTF: return None, 'striprtf not installed'
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return rtf_to_text(f.read()), None
    except Exception as e:
        return None, f'rtf error: {e}'


def read_epub(file_path):
    if not HAS_EPUB: return None, 'ebooklib not installed'
    try:
        book = epub.read_epub(file_path)
        parts = []
        meta = []
        t = book.get_metadata('DC', 'title')
        if t: meta.append(f'标题: {t[0][0]}')
        a = book.get_metadata('DC', 'creator')
        if a: meta.append(f'作者: {a[0][0]}')
        if meta: parts.append('[EPUB 元信息]\n' + '\n'.join(meta))
        texts = []
        for item in book.get_items():
            if item.get_type() == ebooklib.ITEM_DOCUMENT:
                soup = BeautifulSoup(item.get_content(), 'html.parser')
                text = soup.get_text(separator='\n', strip=True)
                if text: texts.append(text)
        parts.append('\n\n'.join(texts))
        return '\n\n'.join(parts), None
    except Exception as e:
        return None, f'epub error: {e}'


def read_odt(file_path):
    try:
        with zipfile.ZipFile(file_path, 'r') as z:
            content = z.open('content.xml').read().decode('utf-8')
        texts = re.findall(r'<text:p[^>]*>(.*?)</text:p>', content, re.DOTALL)
        result = []
        for t in texts:
            t = re.sub(r'<[^>]+>', '', t)
            if t.strip(): result.append(t)
        return '\n\n'.join(result), None
    except Exception as e:
        return None, f'odt error: {e}'


def read_ods(file_path):
    try:
        with zipfile.ZipFile(file_path, 'r') as z:
            content = z.open('content.xml').read().decode('utf-8')
        tables = re.findall(r'<table:table[^>]*>(.*?)</table:table>', content, re.DOTALL)
        parts = []
        for ti, table in enumerate(tables, 1):
            rows = re.findall(r'<table:table-row[^>]*>(.*?)</table:table-row>', table, re.DOTALL)
            row_texts = []
            for row in rows:
                cells = re.findall(r'<table:table-cell[^>]*>(.*?)</table:table-cell>', row, re.DOTALL)
                cell_texts = [re.sub(r'<[^>]+>', '', c) for c in cells]
                if any(c.strip() for c in cell_texts): row_texts.append(' | '.join(cell_texts))
            if row_texts: parts.append(f'=== 表格 {ti} ({len(row_texts)} 行) ===\n' + '\n'.join(row_texts))
        return '\n\n'.join(parts), None
    except Exception as e:
        return None, f'ods error: {e}'


def read_odp(file_path):
    try:
        with zipfile.ZipFile(file_path, 'r') as z:
            content = z.open('content.xml').read().decode('utf-8')
        frames = re.findall(r'<draw:frame[^>]*>(.*?)</draw:frame>', content, re.DOTALL)
        texts = []
        for frame in frames:
            ps = re.findall(r'<text:p[^>]*>(.*?)</text:p>', frame, re.DOTALL)
            for p in ps:
                p = re.sub(r'<[^>]+>', '', p)
                if p.strip(): texts.append(p)
        pages = re.findall(r'<draw:page[^>]*draw:name="([^"]*)"', content)
        result = []
        if pages: result.append('[幻灯片: ' + ', '.join(pages) + ']')
        result.extend(texts)
        return '\n\n'.join(result), None
    except Exception as e:
        return None, f'odp error: {e}'


def read_ole2(file_path):
    if not HAS_OLE: return None, 'olefile not installed'
    try:
        ole = olefile.OleFileIO(file_path)
        streams = ole.listdir()
        if ['WordDocument'] in streams:
            return '[.doc 旧版 Word 格式，建议转换为 .docx 后读取]', None
        if ['Workbook'] in streams:
            if HAS_XLRD:
                try:
                    book = xlrd.open_workbook(file_path)
                    parts = []
                    for si in range(book.nsheets):
                        sheet = book.sheet_by_index(si)
                        rows = []
                        for ri in range(sheet.nrows):
                            vals = [str(sheet.cell_value(ri, ci)) for ci in range(sheet.ncols)]
                            if any(v.strip() for v in vals): rows.append(' | '.join(vals))
                        parts.append(f'=== 工作表: {sheet.name} ({len(rows)} 行) ===\n' + '\n'.join(rows))
                    return '\n\n'.join(parts), None
                except Exception as e:
                    return None, f'xls error: {e}'
            else:
                return '[.xls 旧版 Excel，xlrd 未安装，建议转换为 .xlsx]', None
        if ['PowerPoint Document'] in streams:
            return '[.ppt 旧版 PowerPoint，建议转换为 .pptx 后读取]', None
        stream_names = ['/'.join(s) for s in streams[:10]]
        return '[OLE2 格式，流: ' + ', '.join(stream_names) + ']', None
    except Exception as e:
        return None, f'ole2 error: {e}'


def read_zip(file_path):
    try:
        with zipfile.ZipFile(file_path, 'r') as z:
            entries = z.namelist()
            parts = [f'[ZIP 压缩包: {len(entries)} 个文件/目录]\n']
            readable = {'.txt','.md','.json','.yaml','.yml','.xml','.csv','.js','.ts','.py','.html','.css','.log'}
            text_entries = []
            for name in entries:
                info = z.getinfo(name)
                is_dir = name.endswith('/')
                size_str = f'{info.file_size} B' if info.file_size < 1024 else f'{info.file_size/1024:.1f} KB'
                parts.append(f'  {name}{"/" if is_dir else ""} ({size_str})')
                if not is_dir:
                    ext = Path(name).suffix.lower()
                    if ext in readable and info.file_size < 100000: text_entries.append(name)
            if text_entries and len(text_entries) <= 10:
                parts.append('\n[可读文件内容:]')
                for name in text_entries:
                    try:
                        content = z.open(name).read().decode('utf-8', errors='ignore')
                        parts.append(f'\n--- {name} ---\n{content[:2000]}{"..." if len(content)>2000 else ""}')
                    except:
                        pass
            return '\n'.join(parts), None
    except Exception as e:
        return None, f'zip error: {e}'


def read_file(file_path):
    fmt = detect_format(file_path)
    readers = {
        'docx': read_docx, 'xlsx': read_xlsx, 'pptx': read_pptx,
        'wps': read_docx, 'et': read_xlsx, 'dps': read_pptx,
        'pdf': read_pdf, 'rtf': read_rtf, 'epub': read_epub,
        'odt': read_odt, 'ods': read_ods, 'odp': read_odp,
        'doc': read_ole2, 'xls': read_ole2, 'ppt': read_ole2,
        'ole2': read_ole2, 'zip': read_zip,
    }
    reader = readers.get(fmt)
    if not reader:
        return {'success': False, 'error': f'不支持的文件格式: {fmt}', 'format': fmt}
    text, error = reader(file_path)
    if error:
        return {'success': False, 'error': error, 'format': fmt}
    return {'success': True, 'text': text, 'format': fmt}


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(json.dumps({'success': False, 'error': '缺少文件路径参数'}))
        sys.exit(1)
    file_path = sys.argv[1]
    if not os.path.exists(file_path):
        print(json.dumps({'success': False, 'error': f'文件不存在: {file_path}'}))
        sys.exit(1)
    result = read_file(file_path)
    print(json.dumps(result, ensure_ascii=False))
