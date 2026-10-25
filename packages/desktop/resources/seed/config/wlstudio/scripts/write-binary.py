#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import json
import sys
import shutil
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

ALIGNMENT_MAP = {
    'left': WD_ALIGN_PARAGRAPH.LEFT,
    'center': WD_ALIGN_PARAGRAPH.CENTER,
    'right': WD_ALIGN_PARAGRAPH.RIGHT,
    'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
}

def set_run_font(run, font_name=None, font_size=None, bold=None, italic=None, color=None):
    if font_name:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if font_size is not None:
        run.font.size = Pt(font_size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        color = color.lstrip('#')
        if len(color) == 6:
            run.font.color.rgb = RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))

def op_replace_text(doc, params):
    search = params.get('search', '')
    replace = params.get('replace', '')
    match_case = params.get('matchCase', True)
    count = 0
    if not search:
        return {'success': True, 'replaced': 0}
    targets = []
    for para in doc.paragraphs:
        targets.append(para)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    targets.append(para)
    for para in targets:
        full_text = para.text
        if not match_case:
            if search.lower() in full_text.lower():
                para.clear()
                new_text = full_text.lower().replace(search.lower(), replace)
                run = para.add_run(new_text)
                count += 1
        else:
            if search in full_text:
                para.clear()
                new_text = full_text.replace(search, replace)
                run = para.add_run(new_text)
                count += 1
    return {'success': True, 'replaced': count}

def op_set_style(doc, params):
    p_idx = params.get('paragraphIndex')
    if p_idx is None or p_idx < 0 or p_idx >= len(doc.paragraphs):
        return {'success': False, 'message': 'paragraphIndex invalid'}
    para = doc.paragraphs[p_idx]
    alignment = params.get('alignment')
    if alignment:
        para.alignment = ALIGNMENT_MAP.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)
    for run in para.runs:
        set_run_font(run,
            font_name=params.get('fontName'),
            font_size=params.get('fontSize'),
            bold=params.get('bold'),
            italic=params.get('italic'),
            color=params.get('color'),
        )
    return {'success': True}

def op_add_paragraph(doc, params):
    text = params.get('text', '')
    style = params.get('style')
    alignment = params.get('alignment')
    para = doc.add_paragraph(text, style=style)
    if alignment:
        para.alignment = ALIGNMENT_MAP.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)
    for run in para.runs:
        set_run_font(run,
            font_name=params.get('fontName'),
            font_size=params.get('fontSize'),
            bold=params.get('bold'),
            italic=params.get('italic'),
            color=params.get('color'),
        )
    return {'success': True, 'index': len(doc.paragraphs) - 1}

def op_delete_paragraph(doc, params):
    p_idx = params.get('paragraphIndex')
    if p_idx is None or p_idx < 0 or p_idx >= len(doc.paragraphs):
        return {'success': False, 'message': 'paragraphIndex invalid'}
    p = doc.paragraphs[p_idx]
    p._element.getparent().remove(p._element)
    return {'success': True}

def op_insert_paragraph(doc, params):
    idx = params.get('index', 0)
    text = params.get('text', '')
    alignment = params.get('alignment')
    if idx < 0 or idx > len(doc.paragraphs):
        idx = len(doc.paragraphs)
    if idx == len(doc.paragraphs):
        para = doc.add_paragraph(text)
    else:
        ref = doc.paragraphs[idx]
        para = ref.insert_paragraph_before(text)
    if alignment:
        para.alignment = ALIGNMENT_MAP.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)
    for run in para.runs:
        set_run_font(run,
            font_name=params.get('fontName'),
            font_size=params.get('fontSize'),
            bold=params.get('bold'),
            italic=params.get('italic'),
            color=params.get('color'),
        )
    return {'success': True, 'index': idx}

def op_set_cell(wb, params):
    if not HAS_OPENPYXL:
        return {'success': False, 'message': 'openpyxl not installed'}
    sheet = params.get('sheet', 'Sheet1')
    cell_ref = params.get('cell', 'A1')
    value = params.get('value', '')
    if sheet not in wb.sheetnames:
        return {'success': False, 'message': f'Sheet {sheet} not found'}
    wb[sheet][cell_ref] = value
    return {'success': True}

def op_set_cell_style(wb, params):
    if not HAS_OPENPYXL:
        return {'success': False, 'message': 'openpyxl not installed'}
    sheet = params.get('sheet', 'Sheet1')
    cell_ref = params.get('cell', 'A1')
    if sheet not in wb.sheetnames:
        return {'success': False, 'message': f'Sheet {sheet} not found'}
    cell = wb[sheet][cell_ref]
    fk = {}
    if 'fontName' in params: fk['name'] = params['fontName']
    if 'fontSize' in params: fk['size'] = params['fontSize']
    if 'bold' in params: fk['bold'] = params['bold']
    if 'italic' in params: fk['italic'] = params['italic']
    if 'color' in params:
        c = params['color'].lstrip('#')
        if len(c) == 6: fk['color'] = c
    if fk: cell.font = Font(**fk)
    if 'alignment' in params:
        cell.alignment = Alignment(horizontal=params['alignment'])
    if 'fillColor' in params:
        fc = params['fillColor'].lstrip('#')
        if len(fc) == 6:
            cell.fill = PatternFill(start_color=fc, end_color=fc, fill_type='solid')
    return {'success': True}

DOCX_OPS = {
    'replaceText': op_replace_text,
    'setStyle': op_set_style,
    'addParagraph': op_add_paragraph,
    'deleteParagraph': op_delete_paragraph,
    'insertParagraph': op_insert_paragraph,
}

XLSX_OPS = {
    'setCell': op_set_cell,
    'setCellStyle': op_set_cell_style,
}

def process_docx(config):
    file_path = config['filePath']
    backup = config.get('backup', True)
    if backup:
        shutil.copy2(file_path, file_path + '.bak')
    doc = Document(file_path)
    results = []
    for op in config.get('operations', []):
        handler = DOCX_OPS.get(op.get('type'))
        if not handler:
            results.append({'type': op['type'], 'success': False, 'message': 'unsupported'})
            continue
        try:
            results.append({'type': op['type'], **handler(doc, op.get('params', {}))})
        except Exception as e:
            results.append({'type': op['type'], 'success': False, 'message': str(e)})
    doc.save(file_path)
    return {'success': True, 'results': results}

def process_xlsx(config):
    if not HAS_OPENPYXL:
        return {'success': False, 'message': 'openpyxl not installed'}
    file_path = config['filePath']
    backup = config.get('backup', True)
    if backup:
        shutil.copy2(file_path, file_path + '.bak')
    wb = load_workbook(file_path)
    results = []
    for op in config.get('operations', []):
        handler = XLSX_OPS.get(op.get('type'))
        if not handler:
            results.append({'type': op['type'], 'success': False, 'message': 'unsupported'})
            continue
        try:
            results.append({'type': op['type'], **handler(wb, op.get('params', {}))})
        except Exception as e:
            results.append({'type': op['type'], 'success': False, 'message': str(e)})
    wb.save(file_path)
    return {'success': True, 'results': results}

def main():
    if len(sys.argv) < 2:
        print(json.dumps({'success': False, 'message': 'Usage: write-binary.py <config.json>'}, ensure_ascii=False))
        sys.exit(1)
    with open(sys.argv[1], 'r', encoding='utf-8-sig') as f:
        config = json.load(f)
    ext = config['filePath'].split('.')[-1].lower()
    if ext == 'docx':
        result = process_docx(config)
    elif ext == 'xlsx':
        result = process_xlsx(config)
    else:
        result = {'success': False, 'message': f'Unsupported extension: {ext}'}
    print(json.dumps(result, ensure_ascii=False))

if __name__ == '__main__':
    main()
